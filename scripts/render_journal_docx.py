#!/usr/bin/env python3
"""Render structured project experience JSON to a Word document."""

from __future__ import annotations

import json
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


DATE_PATTERN = re.compile(r"^\d{4}-\d{2}-\d{2}$")
DEFAULT_TITLE = "项目实践与能力成长记录"


def fail(message: str) -> None:
    raise ValueError(message)


def as_items(value: object) -> list[str]:
    if value is None:
        return []
    if isinstance(value, str):
        return [value.strip()] if value.strip() else []
    if isinstance(value, list):
        items = []
        for item in value:
            if not isinstance(item, str):
                fail("Section list items must be strings")
            if item.strip():
                items.append(item.strip())
        return items
    fail("Section values must be strings or lists of strings")


def validate(payload: object) -> dict:
    if not isinstance(payload, dict):
        fail("Input root must be a JSON object")
    entries = payload.get("entries")
    if not isinstance(entries, list) or not entries:
        fail("entries must be a non-empty array")
    for index, entry in enumerate(entries, start=1):
        if not isinstance(entry, dict):
            fail(f"entry {index} must be an object")
        date = entry.get("date")
        if not isinstance(date, str) or not DATE_PATTERN.fullmatch(date):
            fail(f"entry {index} date must use YYYY-MM-DD")
        project = entry.get("project")
        if not isinstance(project, str) or not project.strip():
            fail(f"entry {index} project is required")
        sections = entry.get("sections")
        if not isinstance(sections, dict) or not sections:
            fail(f"entry {index} sections must be a non-empty object")
        for heading, value in sections.items():
            if not isinstance(heading, str) or not heading.strip():
                fail(f"entry {index} contains an invalid section heading")
            as_items(value)
    return payload


def set_font(run, latin: str = "Aptos", east_asia: str = "等线", size: int = 11) -> None:
    run.font.name = latin
    run._element.rPr.rFonts.set(qn("w:eastAsia"), east_asia)
    run.font.size = Pt(size)


def add_items(document: Document, value: object) -> None:
    for item in as_items(value):
        paragraph = document.add_paragraph(style="List Bullet")
        run = paragraph.add_run(item)
        set_font(run)


def render(payload: dict, destination: Path) -> None:
    document = Document()
    section = document.sections[0]
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    section.left_margin = Cm(2.7)
    section.right_margin = Cm(2.7)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title.add_run(str(payload.get("title") or DEFAULT_TITLE))
    set_font(title_run, east_asia="黑体", size=20)
    title_run.bold = True

    period_summary = payload.get("period_summary")
    if isinstance(period_summary, dict) and period_summary:
        document.add_heading("阶段总结", level=1)
        for heading, value in period_summary.items():
            items = as_items(value)
            if not items:
                continue
            document.add_heading(str(heading), level=2)
            add_items(document, items)

    for entry in payload["entries"]:
        stage = str(entry.get("stage") or "").strip()
        heading = f"{entry['date']}｜{entry['project'].strip()}"
        if stage:
            heading += f"｜{stage}"
        document.add_heading(heading, level=1)
        for section_heading, value in entry["sections"].items():
            items = as_items(value)
            if not items:
                continue
            document.add_heading(section_heading.strip(), level=2)
            add_items(document, items)

    styles = document.styles
    for style_name in ("Normal", "List Bullet"):
        style = styles[style_name]
        style.font.name = "Aptos"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "等线")
        style.font.size = Pt(11)
    for style_name, size in (("Title", 20), ("Heading 1", 15), ("Heading 2", 12)):
        style = styles[style_name]
        style.font.name = "Aptos Display"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
        style.font.size = Pt(size)

    destination.parent.mkdir(parents=True, exist_ok=True)
    document.save(destination)


def main() -> int:
    if len(sys.argv) != 3:
        print("Usage: render_journal_docx.py INPUT.json OUTPUT.docx", file=sys.stderr)
        return 2
    source = Path(sys.argv[1])
    destination = Path(sys.argv[2])
    try:
        payload = validate(json.loads(source.read_text(encoding="utf-8")))
        render(payload, destination)
    except (OSError, json.JSONDecodeError, ValueError) as exc:
        print(f"Error: {exc}", file=sys.stderr)
        return 1
    print(destination)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

