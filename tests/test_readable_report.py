import copy
import json
from pathlib import Path
import sys
import tempfile
import unittest

from docx import Document

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / 'scripts'))
from report_model import markdown, validate_report
from render_report import publish_report
from render_journal_docx import render, validate
from lint_report import lint_report
from validate_evidence import preflight


def sample():
    return json.loads((ROOT / 'examples/example-report-v2.json').read_text(encoding='utf-8'))


class ReadableReportTests(unittest.TestCase):
    def test_full_report_has_one_title_and_clean_spacing(self):
        result = markdown(sample())
        self.assertEqual(lint_report(result), ([], []))
        self.assertEqual(sum(line.startswith('# ') for line in result.splitlines()), 1)

    def test_duplicate_section_id_rejected(self):
        data = sample()
        data['sections'][1]['id'] = 'overview'
        with self.assertRaises(ValueError):
            validate_report(data)

    def test_non_object_and_mixed_schema_rejected(self):
        with self.assertRaises(ValueError):
            validate_report([])
        data = sample()
        data['entries'] = []
        with self.assertRaises(ValueError):
            validate_report(data)

    def test_lists_need_spacing(self):
        errors, _ = lint_report('# 标题\n\n正文\n- 条目\n下一段\n')
        self.assertTrue(any('before list' in error for error in errors))
        self.assertTrue(any('after list' in error for error in errors))

    def test_missing_full_section_rejected_brief_allowed(self):
        data = sample()
        data['sections'].pop()
        with self.assertRaises(ValueError):
            validate(data)
        data['profile'] = 'brief'
        self.assertIs(validate(data), data)

    def test_invalid_date_and_multiline_prose_rejected(self):
        for field, value in [('as_of', '2026-02-30'), ('title', 'Title\n# old')]:
            data = sample()
            data[field] = value
            with self.assertRaises(ValueError):
                validate(data)

    def test_broken_blocks_rejected(self):
        for block in ({'type': 'paragraph', 'text': ''}, {'type': 'bullets', 'items': []}, {'type': 'html', 'text': 'x'}):
            data = sample()
            data['sections'][0]['blocks'] = [block]
            with self.assertRaises(ValueError):
                validate(data)

    def test_updates_replace_current_text_and_snapshot_old(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            (root / 'ISSUE_LEDGER.md').write_text('TEST_VERIFIED; not rerun', encoding='utf-8')
            first = sample()
            publish_report(first, root)
            old = (root / 'PROJECT_REPORT.md').read_text(encoding='utf-8')
            changed = copy.deepcopy(first)
            changed['sections'][0]['blocks'][0]['text'] = '新的项目概况，原阶段内容保留在独立快照。'
            publish_report(changed, root)
            current = (root / 'PROJECT_REPORT.md').read_text(encoding='utf-8')
            self.assertNotIn(first['sections'][0]['blocks'][0]['text'], current)
            self.assertEqual(current.count('## 项目概况'), 1)
            snapshots = list((root / 'snapshots').glob('*.md'))
            self.assertEqual(len(snapshots), 1)
            self.assertEqual(snapshots[0].read_text(encoding='utf-8'), old)
            self.assertEqual((root / 'ISSUE_LEDGER.md').read_text(), 'TEST_VERIFIED; not rerun')
            self.assertEqual(changed['sections'][1:], first['sections'][1:])
            publish_report(changed, root)
            self.assertEqual(len(list((root / 'snapshots').iterdir())), 2)

    def test_invalid_update_leaves_current_and_snapshots_untouched(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            publish_report(sample(), root)
            before = (root / 'PROJECT_REPORT.md').read_bytes()
            data = sample()
            data['title'] = '# extra title'
            data['sections'][0]['blocks'][0]['text'] = '# Another title'
            with self.assertRaises(ValueError):
                publish_report(data, root)
            self.assertEqual((root / 'PROJECT_REPORT.md').read_bytes(), before)
            self.assertFalse((root / 'snapshots').exists())

    def test_old_report_concatenation_is_rejected(self):
        content = markdown(sample())
        errors, _ = lint_report(content + '\n# 历史全文\n\n' + content)
        self.assertIn('Report must have exactly one document title', errors)

    def test_heading_jump_and_missing_spacing_rejected(self):
        errors, _ = lint_report('# 标题\n正文\n\n### 跳级\n')
        self.assertTrue(any('blank line' in error for error in errors))
        self.assertTrue(any('level' in error for error in errors))

    def test_duplicate_paragraph_and_long_prose_detected(self):
        paragraph = '重复内容必须被识别。' * 12
        errors, warnings = lint_report('# 标题\n\n' + paragraph + '\n\n' + paragraph + '\n\n' + '长' * 501)
        self.assertIn('Duplicate long paragraph', errors)
        self.assertTrue(warnings)

    def test_code_fence_heading_is_not_counted(self):
        self.assertEqual(lint_report('# 标题\n\n```text\n# code\n```\n'), ([], []))

    def test_v2_privacy_scan_includes_all_prose(self):
        data = sample()
        data['sections'][4]['blocks'][0]['text'] = '联系 reader@example.com'
        self.assertTrue(preflight(data)[0])

    def test_legacy_summary_is_also_scanned(self):
        data = {'entries': [{'date': '2026-10-28', 'project': '示例', 'sections': {'结果': '完成'}}],
                'period_summary': {'概况': '联系 reader@example.com'}}
        self.assertTrue(preflight(validate(data))[0])

    def test_legacy_optional_null_summary_remains_compatible(self):
        data = {'entries': [{'date': '2026-10-28', 'project': '示例', 'sections': {'结果': '完成'}}],
                'period_summary': None}
        self.assertEqual(preflight(validate(data))[0], [])

    def test_word_uses_real_title_normal_prose_and_heading_keep(self):
        with tempfile.TemporaryDirectory() as temp:
            target = Path(temp) / 'report.docx'
            render(sample(), target)
            doc = Document(target)
            self.assertEqual(doc.paragraphs[0].style.name, 'Title')
            normal = [p for p in doc.paragraphs if p.style.name == 'Normal' and p.text]
            self.assertGreaterEqual(len(normal), 5)
            self.assertEqual(len([p for p in doc.paragraphs if p.style.name == 'List Bullet']), 2)
            self.assertTrue(doc.styles['Heading 1'].paragraph_format.keep_with_next)
            self.assertEqual(str(doc.styles['Heading 1'].font.color.rgb), '000000')


if __name__ == '__main__':
    unittest.main()
